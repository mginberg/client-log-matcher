
import streamlit as st
import pandas as pd
import io
from openpyxl import Workbook
from openpyxl.styles import Border, Side, Alignment
from openpyxl.utils import get_column_letter

st.set_page_config(layout="wide")
st.title("Client & Call Log Matcher with Stats")

client_file = st.file_uploader("Upload Client List CSV", type=["csv"])
call_log_file = st.file_uploader("Upload Call Log CSV", type=["csv"])

if client_file and call_log_file:
    client_df = pd.read_csv(client_file)
    call_df = pd.read_csv(call_log_file)

    client_df['Phone Number'] = client_df['Phone Number'].astype(str).str.replace(r"\D", "", regex=True).str[-10:]
    call_df['Clean Phone'] = call_df.iloc[:, 5].astype(str).str.replace(r"\D", "", regex=True).str[-10:]

    matches = []
    for _, client_row in client_df.iterrows():
        phone = client_row['Phone Number']
        matches_df = call_df[call_df['Clean Phone'] == phone]
        if not matches_df.empty:
            longest = matches_df.loc[matches_df.iloc[:, 10].astype(str).str.extract(r'(\d+)')[0].astype(float).idxmax()]
            lead_source = longest.iloc[19]  # Column T
            heap_dnc = "X" if "HEAP DNC UPLINE" in str(longest.iloc[27]) else ""
            vendor = longest.iloc[27] if str(lead_source).strip() == "MEDICARE CLOSER QUEUE" else ""
            matches.append([
                client_row['Client Name'],
                phone,
                lead_source,
                vendor,
                heap_dnc
            ])
        else:
            matches.append([
                client_row['Client Name'],
                phone,
                "NEED TO CHECK",
                "",
                ""
            ])

    # Count duplicates and adjust
    df_final = pd.DataFrame(matches, columns=["Client Name", "Phone Number", "Lead Source", "Vendor", "HEAP DNC"])
    df_final['Duplicate'] = df_final.duplicated(subset='Phone Number', keep='first')
    df_final.loc[df_final['Duplicate'], 'Lead Source'] = 'DUPLICATE'
    df_final.drop(columns='Duplicate', inplace=True)

    # Excel export with formatted stats table
    wb = Workbook()
    ws = wb.active
    ws.title = "Matched_Client_Stats"

    # Header
    headers = ["Client Name", "Phone Number", "Lead Source", "Vendor", "HEAP DNC"]
    ws.append(headers)

    # Add data
    for row in df_final.itertuples(index=False):
        ws.append(list(row))

    # Stats Table
    stats_start_col = 7
    stats_headers = ["QUEUE STATS", "CALLS", "ABANDON", "SALES", "CLOSING %"]
    for i, h in enumerate(stats_headers):
        ws.cell(row=2, column=stats_start_col + i, value=h)

    # Count sales
    queues = sorted(set(df_final["Lead Source"]) - {"DUPLICATE", "NEED TO CHECK"})
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    row_offset = 3
    for i, q in enumerate(queues):
        row_num = i + row_offset
        sales = (df_final["Lead Source"] == q).sum()
        ws.cell(row=row_num, column=stats_start_col, value=q)
        ws.cell(row=row_num, column=stats_start_col + 1, value=0)
        ws.cell(row=row_num, column=stats_start_col + 2, value=0)
        ws.cell(row=row_num, column=stats_start_col + 3, value=sales)
        ws.cell(row=row_num, column=stats_start_col + 4, value=f"=IF(H{row_num}=0,"",J{row_num}/H{row_num})")
        for c in range(stats_start_col, stats_start_col + 5):
            ws.cell(row=row_num, column=c).border = thin_border
            ws.cell(row=row_num, column=c).alignment = Alignment(horizontal="center")

    total_row = len(queues) + row_offset
    ws.cell(row=total_row, column=stats_start_col, value="TOTAL")
    ws.cell(row=total_row, column=stats_start_col + 1, value=f"=SUM(H{row_offset}:H{total_row - 1})")
    ws.cell(row=total_row, column=stats_start_col + 2, value=f"=SUM(I{row_offset}:I{total_row - 1})")
    ws.cell(row=total_row, column=stats_start_col + 3, value=f"=SUM(J{row_offset}:J{total_row - 1})")
    ws.cell(row=total_row, column=stats_start_col + 4, value=f"=IF(H{total_row}=0,"",J{total_row}/H{total_row})")
    for c in range(stats_start_col, stats_start_col + 5):
        ws.cell(row=total_row, column=c).border = thin_border
        ws.cell(row=total_row, column=c).alignment = Alignment(horizontal="center")

    # Auto-width
    for col in range(1, stats_start_col + 6):
        max_len = 0
        col_letter = get_column_letter(col)
        for row in ws.iter_rows(min_col=col, max_col=col):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2

    # Stream download
    output = io.BytesIO()
    wb.save(output)
    st.download_button("Download Final Excel", output.getvalue(), file_name="Matched_Client_Stats.xlsx")
